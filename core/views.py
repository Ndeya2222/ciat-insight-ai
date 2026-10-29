from datetime import timedelta
from functools import wraps
from io import BytesIO
from urllib.parse import urlencode
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.core.paginator import Paginator
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from . import services
from .forms import MentorProjectForm, ProjectForm
import csv, datetime, io
from xml.sax.saxutils import escape
from openpyxl import load_workbook
from .models import ImportLog, DOMAINS, STATUS, ChangeLog, ChatMessage, Cohort, Project


def get_role(u):
    if u.is_superuser: return "admin"
    names = set(u.groups.values_list("name", flat=True))
    for g, r in (("Coordonnateur", "coord"), ("Data Analyst", "analyste"), ("Mentor", "mentor")):
        if g in names: return r
    return None


def role_required(*roles):
    def deco(f):
        @wraps(f)
        def w(request, *a, **k):
            if get_role(request.user) not in roles: raise PermissionDenied
            return f(request, *a, **k)
        return login_required(w)
    return deco

ALL = ("admin", "coord", "analyste", "mentor")
EDIT = ("admin", "coord")


def visible(user):
    """Périmètre de données autorisé (appliqué côté serveur)."""
    qs = Project.objects.select_related("cohort", "mentor")
    return qs.filter(mentor=user) if get_role(user) == "mentor" else qs


def ctx_filters(request):
    return {"cohorts": Cohort.objects.all(), "domains": DOMAINS, "statuses": STATUS, "f": request.GET}


@role_required(*ALL)
def dashboard(request):
    qs = services.apply_filters(visible(request.user), request.GET)
    k = services.kpis(qs)
    cards = [("Candidatures reçues", k["candidatures"], ""), ("Projets sélectionnés", k["selectionnes"], ""),
             ("Projets actifs", k["actifs"], ""), ("Prototypes réalisés", k["prototypes"], ""),
             ("Projets terminés", k["termines"], ""), ("Projets en retard", k["retards"], "text-danger"),
             ("Partenariats actifs", k["partenariats"], ""), ("Financements obtenus (XOF)", f"{k['financements']:,}".replace(",", " "), ""),
             ("Taux de sélection", "n/d" if k["taux_selection"] is None else f"{k['taux_selection']} %", ""),
             ("Taux d'achèvement", "n/d" if k["taux_achevement"] is None else f"{k['taux_achevement']} %", "")]
    return render(request, "core/dashboard.html", {**ctx_filters(request), "cards": cards, "k": k,
        "charts": services.chart_data(qs), "late": services.late_qs(qs).select_related("cohort")[:6]})


@role_required(*ALL)
def project_list(request):
    qs = services.apply_filters(visible(request.user), request.GET)
    if request.GET.get("q"):
        from django.db.models import Q
        q = request.GET["q"]
        qs = qs.filter(Q(title__icontains=q) | Q(description__icontains=q) | Q(owner__icontains=q))
    page = Paginator(qs, 10).get_page(request.GET.get("page"))
    keep = urlencode({k: v for k, v in request.GET.items() if k != "page" and v})
    return render(request, "core/project_list.html", {**ctx_filters(request), "page": page, "keep": keep,
        "can_edit": get_role(request.user) in EDIT})


@role_required(*ALL)
def project_detail(request, pk):
    p = get_object_or_404(visible(request.user), pk=pk)
    return render(request, "core/project_detail.html", {"p": p, "milestones": p.milestones.all(),
        "fundings": p.fundings.all(), "logs": p.logs.all()[:15], "can_edit": get_role(request.user) in EDIT or p.mentor_id == request.user.id and get_role(request.user) == "mentor",
        "can_delete": get_role(request.user) in EDIT})


def _save(request, form, creating):
    old = None if creating else Project.objects.get(pk=form.instance.pk)
    p = form.save(commit=False)
    if p.status == "selectionne" and not p.selected_on: p.selected_on = timezone.localdate()
    p.save()
    if creating: ChangeLog.objects.create(project=p, user=request.user, text="Projet créé")
    else:
        for f in form.changed_data:
            if f in ("status", "progress", "deadline", "mentor"):
                ChangeLog.objects.create(project=p, user=request.user, text=f"{p._meta.get_field(f).verbose_name} : {getattr(old, f)} → {getattr(p, f)}")
    return p


@role_required(*EDIT)
def project_create(request):
    form = ProjectForm(request.POST or None, initial={"submitted_on": timezone.localdate()})
    if request.method == "POST" and form.is_valid():
        p = _save(request, form, True)
        messages.success(request, "Projet créé.")
        return redirect("project_detail", pk=p.pk)
    return render(request, "core/project_form.html", {"form": form, "title": "Nouveau projet"})


@role_required(*ALL)
def project_edit(request, pk):
    p = get_object_or_404(visible(request.user), pk=pk)
    role = get_role(request.user)
    if role not in EDIT and role != "mentor": raise PermissionDenied
    form = (ProjectForm if role in EDIT else MentorProjectForm)(request.POST or None, instance=p)
    if request.method == "POST" and form.is_valid():
        _save(request, form, False)
        messages.success(request, "Modifications enregistrées.")
        return redirect("project_detail", pk=p.pk)
    return render(request, "core/project_form.html", {"form": form, "title": f"Modifier : {p.title}"})


@role_required(*EDIT)
@require_POST
def project_delete(request, pk):
    p = get_object_or_404(Project, pk=pk)
    if p.fundings.exists():
        messages.error(request, "Suppression refusée : des financements sont rattachés à ce projet.")
        return redirect("project_detail", pk=pk)
    p.delete()
    messages.success(request, "Projet supprimé.")
    return redirect("project_list")


@role_required(*ALL)
def project_export(request):
    qs = services.apply_filters(visible(request.user), request.GET)
    wb = Workbook(); ws = wb.active; ws.title = "Projets"
    ws.append(["Titre", "Porteur", "Domaine", "Cohorte", "Statut", "Progression %", "Date limite", "Budget XOF", "En retard"])
    st = dict(STATUS); dm = dict(DOMAINS)
    for p in qs:
        ws.append([p.title, p.owner, dm[p.domain], p.cohort.name, st[p.status], p.progress, p.deadline, int(p.budget), "Oui" if p.is_late else "Non"])
    ws2 = wb.create_sheet("Informations")
    ws2.append(["Généré le", timezone.localtime().strftime("%d/%m/%Y %H:%M")])
    ws2.append(["Filtres appliqués", str({k: v for k, v in request.GET.items() if v}) or "aucun"])
    if qs.filter(is_demo=True).exists(): ws2.append(["Mention", "Document de démonstration (données fictives)"])
    buf = BytesIO(); wb.save(buf)
    r = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    r["Content-Disposition"] = 'attachment; filename="projets_ciat.xlsx"'
    return r


@role_required(*ALL)
def assistant(request):
    if request.method == "POST":
        if "purge" in request.POST:
            ChatMessage.objects.filter(user=request.user).delete()
            messages.success(request, "Historique supprimé.")
        else:
            question = request.POST.get("question", "").strip()[:500]
            recent = ChatMessage.objects.filter(user=request.user, created__gte=timezone.now() - timedelta(minutes=1)).count()
            if not question: messages.error(request, "Veuillez saisir une question.")
            elif recent >= 10: messages.error(request, "Trop de questions : réessayez dans une minute.")
            else:
                text, mode = services.answer(question, visible(request.user))
                ChatMessage.objects.create(user=request.user, question=question, answer=text, mode=mode)
        return redirect("assistant")
    return render(request, "core/assistant.html", {"history": ChatMessage.objects.filter(user=request.user)[:10]})


# ---------------- Import CSV / Excel ----------------
COLS = ["titre", "porteur", "domaine", "cohorte", "date_soumission", "budget"]


def _read(f):
    name = f.name.lower()
    if f.size > 2 * 1024 * 1024: raise ValueError("Fichier trop volumineux (2 Mo maximum).")
    if name.endswith(".csv"):
        try: txt = f.read().decode("utf-8-sig")
        except UnicodeDecodeError: raise ValueError("Encodage non reconnu : enregistrez le CSV en UTF-8.")
        first = txt.split("\n")[0]
        rows = list(csv.reader(io.StringIO(txt), delimiter=";" if first.count(";") > first.count(",") else ","))
    elif name.endswith(".xlsx"):
        ws = load_workbook(f, read_only=True, data_only=True).active
        rows = [["" if c is None else c for c in r] for r in ws.iter_rows(values_only=True)]
    else:
        raise ValueError("Format refusé : seuls .csv et .xlsx sont acceptés.")
    if not rows: raise ValueError("Fichier vide.")
    head = [str(h).strip().lower() for h in rows[0]]
    missing = [c for c in COLS if c not in head]
    if missing: raise ValueError("Colonnes manquantes : " + ", ".join(missing))
    return [dict(zip(head, r)) for r in rows[1:] if any(str(x).strip() for x in r)][:500]


def _validate(rows):
    coh = {c.name.lower(): c for c in Cohort.objects.all()}
    dom = {**{k: k for k, _ in DOMAINS}, **{l.lower(): k for k, l in DOMAINS}}
    ok, errs, seen = [], [], set()
    for i, r in enumerate(rows, start=2):
        e = []
        title, owner = str(r.get("titre", "")).strip(), str(r.get("porteur", "")).strip()
        d = dom.get(str(r.get("domaine", "")).strip().lower())
        c = coh.get(str(r.get("cohorte", "")).strip().lower())
        raw = r.get("date_soumission")
        try:
            date = raw.date() if isinstance(raw, datetime.datetime) else raw if isinstance(raw, datetime.date) else datetime.date.fromisoformat(str(raw).strip())
        except ValueError:
            date = None
        try: budget = int(float(str(r.get("budget") or 0).replace(" ", "")))
        except ValueError: budget = -1
        if not title: e.append("titre manquant")
        if not owner: e.append("porteur manquant")
        if not d: e.append("domaine inconnu")
        if not c: e.append("cohorte inconnue")
        if not date: e.append("date invalide (format AAAA-MM-JJ)")
        if budget < 0: e.append("budget invalide")
        if c and title:
            key = (title.lower(), c.pk)
            if key in seen or Project.objects.filter(title__iexact=title, cohort=c).exists():
                e.append("doublon : projet déjà présent, il ne sera pas écrasé")
            seen.add(key)
        if e: errs.append((i, title or "(sans titre)", "; ".join(e)))
        else: ok.append({"titre": title, "porteur": owner, "domaine": d, "cohorte": c.pk, "date": date.isoformat(), "budget": budget})
    return ok, errs


@role_required(*EDIT)
def project_import(request):
    ctx = {"history": ImportLog.objects.all()[:5]}
    if request.method == "POST":
        if "confirm" in request.POST:
            rows = request.session.pop("import_rows", None)
            if not rows:
                messages.error(request, "Aucun aperçu à confirmer : rechargez le fichier.")
                return redirect("project_import")
            made_n = 0
            for r in rows:
                p, made = Project.objects.get_or_create(title=r["titre"], cohort_id=r["cohorte"], defaults=dict(
                    owner=r["porteur"], domain=r["domaine"], submitted_on=datetime.date.fromisoformat(r["date"]), budget=r["budget"]))
                if made:
                    made_n += 1
                    ChangeLog.objects.create(project=p, user=request.user, text="Projet importé")
            ImportLog.objects.create(user=request.user, filename=request.session.pop("import_name", ""), created=made_n,
                                     errors=request.session.pop("import_errors", 0))
            messages.success(request, f"Import terminé : {made_n} projet(s) créé(s).")
            return redirect("project_list")
        f = request.FILES.get("file")
        if not f: messages.error(request, "Choisissez un fichier.")
        else:
            try: rows = _read(f)
            except ValueError as e: messages.error(request, str(e))
            except Exception: messages.error(request, "Fichier illisible ou corrompu.")
            else:
                ok, errs = _validate(rows)
                request.session.update({"import_rows": ok, "import_name": f.name[:200], "import_errors": len(errs)})
                ctx.update(preview=ok[:20], n_ok=len(ok), errs=errs, filename=f.name)
    return render(request, "core/project_import.html", ctx)


# ---------------- Rapport PDF ----------------
@role_required(*ALL)
def report_pdf(request):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    qs = services.apply_filters(visible(request.user), request.GET)
    k = services.kpis(qs)
    _, recos, limits = services.compute_answer("recommandations", qs)
    ss = getSampleStyleSheet()
    P = lambda t, st="BodyText": Paragraph(escape(t), ss[st])
    nd = lambda v: "n/d" if v is None else f"{v} %"
    dates = qs.order_by("submitted_on").values_list("submitted_on", flat=True)
    period = f"soumissions du {dates.first():%d/%m/%Y} au {dates.last():%d/%m/%Y}" if qs.exists() else "aucune donnée"
    flt = ", ".join(f"{a}={b}" for a, b in request.GET.items() if b) or "aucun"
    import os
    from django.contrib.staticfiles import finders
    from reportlab.platypus import Image
    lp = finders.find("core/img/isep-logo.png")
    st = [Image(lp, width=110, height=110 * 367 / 479, hAlign="LEFT")] if lp else []
    st += [P("CIAT Insight AI - Rapport de synthèse", "Title"), P("Des données aux décisions", "Italic"),
          P(f"Généré le {timezone.localtime():%d/%m/%Y à %H:%M} - Période analysée : {period} - Filtres appliqués : {flt}")]
    if qs.filter(is_demo=True).exists(): st.append(P("Document de démonstration : données fictives.", "Heading3"))
    rows = [["Indicateur", "Valeur"], ["Candidatures reçues", k["candidatures"]], ["Projets sélectionnés", k["selectionnes"]],
            ["Projets actifs", k["actifs"]], ["Prototypes réalisés", k["prototypes"]], ["Projets terminés", k["termines"]],
            ["Projets en retard", k["retards"]], ["Partenariats actifs", k["partenariats"]], ["Financements obtenus (XOF)", f"{k['financements']:,}".replace(",", " ")],
            ["Taux de sélection", nd(k["taux_selection"])], ["Taux d'achèvement", nd(k["taux_achevement"])]]
    t = Table(rows, colWidths=[300, 150])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4c311f")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                           ("GRID", (0, 0), (-1, -1), .5, colors.grey), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f6f9")])]))
    st += [Spacer(1, 10), P("Indicateurs", "Heading2"), t, P("Principaux constats", "Heading2")]
    late = list(services.late_qs(qs)[:10])
    st += [P(f"- {p.title} : échéance {p.deadline or 'non définie'}, jalons en retard {p.late_milestones}") for p in late] or [P("Aucun retard détecté.")]
    st.append(P("Recommandations (hypothèses de travail)", "Heading2"))
    st += [P("- " + r) for r in recos] or [P("Aucune recommandation déclenchée par les données actuelles.")]
    st.append(P("Limites des données", "Heading2"))
    st += [P("- " + l) for l in limits + [f"Taux calculés sur {k['decides']} candidature(s) avec décision ; « n/d » = données insuffisantes."]]
    buf = BytesIO(); SimpleDocTemplate(buf, pagesize=A4, title="Rapport CIAT Insight AI").build(st)
    r = HttpResponse(buf.getvalue(), content_type="application/pdf")
    r["Content-Disposition"] = 'attachment; filename="rapport_ciat.pdf"'
    return r
