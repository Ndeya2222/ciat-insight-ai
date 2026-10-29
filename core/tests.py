from datetime import timedelta
from django.contrib.auth.models import User
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from .models import ChatMessage, Cohort, Funding, Milestone, Project
from . import services


class Base(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("seed_demo", verbosity=0)
        cls.admin = User.objects.get(username="admin_demo")
        cls.coord = User.objects.get(username="coord_demo")
        cls.analyste = User.objects.get(username="analyste_demo")
        cls.mentor = User.objects.get(username="mentor_demo")

    def login(self, u): self.client.force_login(u)


class AuthPermTests(Base):
    def test_login_required(self):
        r = self.client.get(reverse("dashboard"))
        self.assertEqual(r.status_code, 302); self.assertIn("/connexion/", r["Location"])

    def test_login_logout(self):
        self.assertTrue(self.client.login(username="coord_demo", password="Demo@CIAT2026"))
        self.assertEqual(self.client.get(reverse("dashboard")).status_code, 200)
        self.client.post(reverse("logout"))
        self.assertEqual(self.client.get(reverse("dashboard")).status_code, 302)

    def test_analyste_cannot_create_or_delete(self):
        self.login(self.analyste)
        self.assertEqual(self.client.get(reverse("project_create")).status_code, 403)
        p = Project.objects.first()
        self.assertEqual(self.client.post(reverse("project_delete", args=[p.pk])).status_code, 403)
        self.assertTrue(Project.objects.filter(pk=p.pk).exists())

    def test_mentor_scope(self):
        self.login(self.mentor)
        mine = Project.objects.filter(mentor=self.mentor); other = Project.objects.exclude(mentor=self.mentor).first()
        self.assertEqual(self.client.get(reverse("project_detail", args=[other.pk])).status_code, 404)
        self.assertEqual(self.client.get(reverse("project_detail", args=[mine.first().pk])).status_code, 200)
        self.assertEqual(self.client.get(reverse("project_create")).status_code, 403)


class ProjectTests(Base):
    def data(self, **kw):
        d = dict(title="Nouveau", owner="Test", domain="tic", cohort=Cohort.objects.first().pk, submitted_on="2026-01-10",
                 status="soumis", progress=0, budget=1000, description="", problem="", solution="", objectives="")
        d.update(kw); return d

    def test_create_ok_and_persisted(self):
        self.login(self.coord)
        r = self.client.post(reverse("project_create"), self.data())
        self.assertEqual(r.status_code, 302); self.assertTrue(Project.objects.filter(title="Nouveau").exists())

    def test_validation_errors(self):
        self.login(self.coord)
        for bad in (dict(progress=150), dict(budget=-5), dict(deadline="2025-01-01"), dict(submitted_on="")):
            self.assertEqual(self.client.post(reverse("project_create"), self.data(**bad)).status_code, 200)
        self.assertFalse(Project.objects.filter(title="Nouveau").exists())

    def test_invalid_status_transition_rejected(self):
        self.login(self.coord)
        p = Project.objects.filter(status="soumis").first()
        r = self.client.post(reverse("project_edit", args=[p.pk]), self.data(title=p.title, status="termine", progress=100, cohort=p.cohort_id, submitted_on=p.submitted_on))
        self.assertEqual(r.status_code, 200); p.refresh_from_db(); self.assertEqual(p.status, "soumis")

    def test_valid_transition_logged(self):
        self.login(self.coord)
        p = Project.objects.filter(status="soumis").first()
        self.client.post(reverse("project_edit", args=[p.pk]), self.data(title=p.title, status="evaluation", cohort=p.cohort_id, submitted_on=p.submitted_on))
        p.refresh_from_db(); self.assertEqual(p.status, "evaluation"); self.assertTrue(p.logs.filter(text__contains="Statut").exists())

    def test_delete_blocked_with_funding(self):
        self.login(self.coord)
        p = Funding.objects.first().project
        self.client.post(reverse("project_delete", args=[p.pk])); self.assertTrue(Project.objects.filter(pk=p.pk).exists())


class KpiTests(Base):
    def test_kpis_match_database(self):
        k = services.kpis(Project.objects.all())
        self.assertEqual(k["candidatures"], Project.objects.count())
        self.assertEqual(k["selectionnes"], Project.objects.filter(selected_on__isnull=False).count())
        self.assertEqual(k["termines"], Project.objects.filter(status="termine").count())
        sel, rej = k["selectionnes"], Project.objects.filter(status="rejete").count()
        self.assertEqual(k["taux_selection"], round(100 * sel / (sel + rej), 1))
        total = sum(f.amount for f in Funding.objects.filter(status="obtenu", currency="XOF"))
        self.assertEqual(k["financements"], int(total))

    def test_empty_data_no_zero_division(self):
        k = services.kpis(Project.objects.none())
        self.assertIsNone(k["taux_selection"]); self.assertIsNone(k["taux_achevement"]); self.assertEqual(k["candidatures"], 0)

    def test_filters_change_results(self):
        c = Cohort.objects.first()
        self.login(self.coord)
        full = self.client.get(reverse("dashboard")).context["k"]["candidatures"]
        part = self.client.get(reverse("dashboard"), {"cohorte": c.pk}).context["k"]["candidatures"]
        self.assertEqual(part, Project.objects.filter(cohort=c).count()); self.assertLess(part, full)

    def test_late_detection(self):
        c = Cohort.objects.first(); t = timezone.localdate()
        ok = Project.objects.create(title="OK", owner="x", domain="tic", cohort=c, submitted_on=t, status="incubation", deadline=t + timedelta(days=30))
        Milestone.objects.create(project=ok, title="M", due_date=t + timedelta(days=5))
        late_ms = Project.objects.create(title="JalonRetard", owner="x", domain="tic", cohort=c, submitted_on=t, status="incubation")
        Milestone.objects.create(project=late_ms, title="M", due_date=t - timedelta(days=3))
        late_dl = Project.objects.create(title="DeadlineDepassee", owner="x", domain="tic", cohort=c, submitted_on=t, status="incubation", deadline=t - timedelta(days=1))
        done = Project.objects.create(title="Termine", owner="x", domain="tic", cohort=c, submitted_on=t, status="termine", deadline=t - timedelta(days=9))
        ids = set(services.late_qs(Project.objects.all()).values_list("title", flat=True))
        self.assertTrue({"JalonRetard", "DeadlineDepassee"} <= ids); self.assertNotIn("OK", ids); self.assertNotIn("Termine", ids)


class ExportAssistantTests(Base):
    def test_excel_export(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self.login(self.analyste)
        r = self.client.get(reverse("project_export"))
        ws = load_workbook(BytesIO(r.content))["Projets"]
        self.assertEqual(ws.max_row - 1, Project.objects.count())

    def test_assistant_demo_mode_is_transparent_and_correct(self):
        self.login(self.coord)
        self.client.post(reverse("assistant"), {"question": "Combien de projets sont en incubation ?"})
        m = ChatMessage.objects.latest("created")
        self.assertEqual(m.mode, "demo"); self.assertIn("aucun modèle d'IA n'a été interrogé", m.answer)
        self.assertIn(f"{Project.objects.filter(status='incubation').count()} projet(s)", m.answer)

    def test_assistant_provider_down_falls_back(self):
        from unittest import mock
        with self.settings(AI_API_KEY="fake"), mock.patch("urllib.request.urlopen", side_effect=OSError):
            text, mode = services.answer("Quels projets sont en retard ?", Project.objects.all())
        self.assertEqual(mode, "demo"); self.assertIn("indisponible", text)

    def test_assistant_history_purge(self):
        self.login(self.coord)
        self.client.post(reverse("assistant"), {"question": "Compare les cohortes"})
        self.client.post(reverse("assistant"), {"purge": "1"})
        self.assertEqual(ChatMessage.objects.filter(user=self.coord).count(), 0)

    def test_seed_idempotent(self):
        n = Project.objects.count(); call_command("seed_demo", verbosity=0); self.assertEqual(Project.objects.count(), n)


class ImportPdfTests(Base):
    def upload(self, content, name="p.csv"):
        from django.core.files.uploadedfile import SimpleUploadedFile
        return self.client.post(reverse("project_import"), {"file": SimpleUploadedFile(name, content.encode())})

    CSV = ("titre,porteur,domaine,cohorte,date_soumission,budget\nNouveauA,Awa,tic,Cohorte 2026-2027 (fictive),2026-10-01,1000\n"
           "Mauvais,,xyz,Inconnue,nope,-1\n")

    def test_preview_does_not_save_then_confirm_saves(self):
        self.login(self.coord); n = Project.objects.count()
        r = self.upload(self.CSV)
        self.assertEqual(r.context["n_ok"], 1); self.assertEqual(len(r.context["errs"]), 1)
        self.assertEqual(Project.objects.count(), n)
        self.client.post(reverse("project_import"), {"confirm": "1"})
        self.assertEqual(Project.objects.count(), n + 1)
        from .models import ImportLog; self.assertEqual(ImportLog.objects.count(), 1)

    def test_duplicates_never_overwrite(self):
        self.login(self.coord); p = Project.objects.first()
        csv_ = f"titre,porteur,domaine,cohorte,date_soumission,budget\n{p.title},Autre,tic,{p.cohort.name},2026-10-01,1\n"
        r = self.upload(csv_); self.assertEqual(r.context["n_ok"], 0)
        self.assertIn("doublon", r.context["errs"][0][2]); p.refresh_from_db(); self.assertNotEqual(p.owner, "Autre")

    def test_bad_files_rejected_and_permission(self):
        self.login(self.coord)
        self.assertContains(self.upload("x", "virus.exe"), "Format refusé")
        self.assertContains(self.upload("a,b\n1,2\n"), "Colonnes manquantes")
        self.login(self.analyste); self.assertEqual(self.client.get(reverse("project_import")).status_code, 403)

    def test_pdf_report(self):
        self.login(self.analyste)
        r = self.client.get(reverse("report_pdf"), {"cohorte": Cohort.objects.first().pk})
        self.assertTrue(r.content.startswith(b"%PDF")); self.assertEqual(r["Content-Type"], "application/pdf")
